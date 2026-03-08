import os
from typing import Any, Dict, List

try:
    from datasets import load_dataset  # type: ignore
except Exception:
    load_dataset = None  # type: ignore

try:
    from openpyxl import Workbook  # type: ignore
except Exception:
    Workbook = None  # type: ignore


def safe_get(d: Dict[str, Any], path: List[str], default: Any = "") -> Any:
    cur: Any = d
    for p in path:
        if isinstance(cur, dict):
            cur = cur.get(p)
        else:
            return default
    return cur if cur is not None else default


def export_dataset_xlsx(
    dataset_name: str = "Logistic12/weiboDataWithCommentByTheme",
    output_xlsx: str = "output/raw_weibo.xlsx",
) -> None:
    if load_dataset is None:
        raise RuntimeError("缺少 datasets，请先安装：pip install datasets")
    if Workbook is None:
        raise RuntimeError("缺少 openpyxl，请先安装：pip install openpyxl")
    ds = load_dataset(dataset_name)
    os.makedirs(os.path.dirname(output_xlsx), exist_ok=True)
    wb = Workbook()
    ws_articles = wb.active
    ws_articles.title = "articles"
    ws_comments = wb.create_sheet("comments")

    ws_articles.append(
        [
            "keyword",
            "weibo_id",
            "text",
            "author",
            "up",
            "re",
            "cm",
            "comments_count",
            "cleaned_comments_count",
        ]
    )
    ws_comments.append(
        [
            # article fields
            "keyword",
            "weibo_id",
            "article_text",
            "author",
            "up",
            "re",
            "cm",
            # comment fields
            "original_post_id",
            "comment_id",
            "user",
            "content",
            "likes",
            # flags placeholders (raw没有AI标记，保持一致列)
            "ai_ad_spam",
            "ai_meaningless",
            "ai_sensitive",
            "machine_water_army",
            "truncated",
            "sensitive_hits",
        ]
    )

    for split_name in ds.keys():
        split = ds[split_name]
        for r in split:
            r = dict(r)
            details = r.get("weibo_details") or {}
            stats = details.get("stats") or {}
            top_comments = r.get("top_comments") or []
            cleaned_comments = r.get("cleaned_top_comments") or []
            ws_articles.append(
                [
                    r.get("keyword", ""),
                    details.get("id", ""),
                    details.get("text", ""),
                    details.get("author", ""),
                    safe_get(stats, ["up"], 0),
                    safe_get(stats, ["re"], 0),
                    safe_get(stats, ["cm"], 0),
                    len(top_comments),
                    len(cleaned_comments) if cleaned_comments else 0,
                ]
            )
            for c in top_comments:
                ws_comments.append(
                    [
                        r.get("keyword", ""),
                        details.get("id", ""),
                        details.get("text", ""),
                        details.get("author", ""),
                        safe_get(stats, ["up"], 0),
                        safe_get(stats, ["re"], 0),
                        safe_get(stats, ["cm"], 0),
                        c.get("original_post_id", ""),
                        c.get("comment_id", ""),
                        c.get("user", ""),
                        c.get("content", ""),
                        c.get("likes", 0),
                        False,
                        False,
                        False,
                        False,
                        False,
                        "",
                    ]
                )
    wb.save(output_xlsx)


def main() -> None:
    export_dataset_xlsx()
    print("已生成 output/raw_weibo.xlsx")


if __name__ == "__main__":
    main()
